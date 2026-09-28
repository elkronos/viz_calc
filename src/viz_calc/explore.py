"""First-look profiling of a whole DataFrame as pages of small multiples."""

from __future__ import annotations

from collections.abc import Iterable, Sequence
from io import BytesIO
from itertools import combinations
from typing import Any

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from . import stats as st
from ._core import NEUTRAL, OKABE_ITO, VizResult, check_dataframe, require

__all__ = ["profile_bars", "profile_boxes", "profile_scatters", "to_pptx"]


def _split(data: pd.DataFrame, max_levels: int) -> tuple[list[str], list[str]]:
    """Columns treated as categorical (few levels) and as continuous numeric."""
    cat, num = [], []
    for c in data.columns:
        s = data[c]
        is_num = pd.api.types.is_numeric_dtype(s) and not pd.api.types.is_bool_dtype(s)
        if is_num and s.nunique() > max_levels:
            num.append(c)
        elif s.nunique() <= max_levels:
            cat.append(c)
    return cat, num


def _pages(n: int, per_page: int, ncols: int):
    for start in range(0, n, per_page):
        count = min(per_page, n - start)
        nc = min(ncols, count)
        nr = int(np.ceil(count / nc))
        fig, axes = plt.subplots(nr, nc, figsize=(4 * nc, 3.2 * nr), squeeze=False)
        for ax in axes.flat[count:]:
            ax.set_visible(False)
        yield start, fig, list(axes.flat[:count])


def profile_bars(
    data: pd.DataFrame,
    columns: Sequence[str] | None = None,
    max_levels: int = 10,
    labels: bool = True,
    per_page: int = 12,
    ncols: int = 4,
    color: str = OKABE_ITO[0],
) -> VizResult:
    """Bar chart of value counts for every low-cardinality column.

    A column is included when it has at most *max_levels* distinct values.
    Missing values are counted as their own bar so they are not hidden.
    Charts are laid out *per_page* to a figure; ``figure`` is the list of
    figures and ``table`` has the counts.
    """
    check_dataframe(data, columns or [])
    cols = list(columns) if columns is not None else [c for c in data.columns if data[c].nunique() <= max_levels]
    if not cols:
        raise ValueError(f"no columns with at most {max_levels} distinct values")
    figs, rows = [], []
    for start, fig, axes in _pages(len(cols), per_page, ncols):
        for ax, c in zip(axes, cols[start:]):
            counts = data[c].value_counts(dropna=False, sort=False)
            try:
                counts = counts.sort_index()
            except TypeError:
                pass
            names = ["(missing)" if pd.isna(k) else str(k) for k in counts.index]
            bars = ax.bar(range(len(counts)), counts.to_numpy(), color=color)
            if labels:
                ax.bar_label(bars, fontsize=7, padding=1)
            ax.set_xticks(range(len(counts)), names, rotation=45 if max(map(len, names)) > 3 else 0, ha="right" if max(map(len, names)) > 3 else "center")
            ax.set_title(f"{c}  (n={int(counts.sum())})", fontsize="medium", loc="left")
            ax.spines[["top", "right"]].set_visible(False)
            rows += [{"column": c, "level": n, "count": int(v)} for n, v in zip(names, counts)]
        fig.tight_layout()
        figs.append(fig)
    return VizResult(figs, None, pd.DataFrame(rows))


def profile_boxes(
    data: pd.DataFrame,
    max_levels: int = 10,
    categorical: Sequence[str] | None = None,
    numeric: Sequence[str] | None = None,
    per_page: int = 12,
    ncols: int = 4,
) -> VizResult:
    """Box plot of every numeric column by every categorical column.

    Categorical columns have at most *max_levels* distinct values; numeric
    columns have more. The table ranks pairs by eta-squared (the share of the
    numeric column's variance explained by the grouping), so the strongest
    associations can be looked at first.
    """
    cat, num = _split(data, max_levels)
    cat = list(categorical) if categorical is not None else cat
    num = list(numeric) if numeric is not None else num
    check_dataframe(data, [*cat, *num])
    pairs = [(c, n) for c in cat for n in num if c != n]
    if not pairs:
        raise ValueError("no (categorical, numeric) column pairs found; adjust max_levels")
    rows = []
    for c, n in pairs:
        d = data[[c, n]].dropna()
        grand = d[n].mean()
        ss_tot = ((d[n] - grand) ** 2).sum()
        g = d.groupby(c, observed=True)[n]
        ss_between = (g.size() * (g.mean() - grand) ** 2).sum()
        rows.append({"categorical": c, "numeric": n, "n": len(d), "eta_squared": ss_between / ss_tot if ss_tot else np.nan})
    table = pd.DataFrame(rows).sort_values("eta_squared", ascending=False, ignore_index=True)
    figs = []
    for start, fig, axes in _pages(len(table), per_page, ncols):
        for ax, r in zip(axes, table.iloc[start:].itertuples()):
            d = data[[r.categorical, r.numeric]].dropna()
            levels = sorted(d[r.categorical].unique(), key=str)
            ax.boxplot([d.loc[d[r.categorical] == lv, r.numeric] for lv in levels], showfliers=True,
                       flierprops={"markersize": 2, "markeredgecolor": NEUTRAL}, medianprops={"color": OKABE_ITO[5]})
            ax.set_xticks(range(1, len(levels) + 1), [str(lv) for lv in levels], rotation=45, ha="right")
            ax.set_title(f"{r.numeric} by {r.categorical}  (η²={r.eta_squared:.2f})", fontsize="small", loc="left")
            ax.spines[["top", "right"]].set_visible(False)
        fig.tight_layout()
        figs.append(fig)
    return VizResult(figs, None, table)


def profile_scatters(
    data: pd.DataFrame,
    columns: Sequence[str] | None = None,
    min_levels: int = 10,
    method: str = "pearson",
    p_adjust: str = "holm",
    fit: bool = True,
    per_page: int = 12,
    ncols: int = 4,
    max_points: int = 5000,
) -> VizResult:
    """Scatter plot for every pair of numeric columns, strongest first.

    Numeric columns with at least *min_levels* distinct values are used.
    Pairs are ordered by absolute correlation, and each title shows *r* with a
    multiplicity-adjusted p-value. Large data are subsampled to *max_points*
    points per panel for drawing only (statistics use all rows).
    """
    if columns is None:
        columns = [c for c in data.columns if pd.api.types.is_numeric_dtype(data[c])
                   and not pd.api.types.is_bool_dtype(data[c]) and data[c].nunique() >= min_levels]
    check_dataframe(data, columns)
    pairs = list(combinations(columns, 2))
    if not pairs:
        raise ValueError(f"need at least two numeric columns with ≥{min_levels} distinct values")
    table = pd.DataFrame([{"x": a, "y": b, **st.correlation_test(data[a], data[b], method)} for a, b in pairs])
    table["p_adjusted"] = st.adjust_pvalues(table["p"], p_adjust)
    table = table.reindex(table["r"].abs().sort_values(ascending=False).index).reset_index(drop=True)
    figs = []
    for start, fig, axes in _pages(len(table), per_page, ncols):
        for ax, r in zip(axes, table.iloc[start:].itertuples()):
            d = data[[r.x, r.y]].dropna()
            if len(d) > max_points:
                d = d.sample(max_points, random_state=0)
            ax.scatter(d[r.x], d[r.y], s=6, alpha=0.5, color=OKABE_ITO[0], lw=0)
            if fit and len(d) > 2 and d[r.x].nunique() > 1:
                slope, icpt = np.polyfit(d[r.x], d[r.y], 1)
                gx = np.linspace(d[r.x].min(), d[r.x].max(), 20)
                ax.plot(gx, icpt + slope * gx, color=OKABE_ITO[5], lw=1.2)
            ax.set_xlabel(r.x, fontsize=8)
            ax.set_ylabel(r.y, fontsize=8)
            ax.set_title(f"r={r.r:.2f}, p_adj={r.p_adjusted:.2g}", fontsize="small", loc="left")
            ax.spines[["top", "right"]].set_visible(False)
        fig.tight_layout()
        figs.append(fig)
    return VizResult(figs, None, table, {"method": method, "p_adjust": p_adjust})


def to_pptx(figures: Any, path: str, titles: Iterable[str] | None = None, dpi: int = 200) -> str:
    """Write figures (or :class:`VizResult` objects) to a PowerPoint file, one per slide.

    Requires ``python-pptx`` (``pip install "viz_calc[pptx]"``). Images are
    scaled to fit a 13.33 × 7.5 inch (16:9) slide while keeping their aspect
    ratio.
    """
    pptx = require("pptx", "pptx")
    from pptx.util import Inches, Pt

    items = figures if isinstance(figures, (list, tuple)) else [figures]
    flat = []
    for it in items:
        fig = it.figure if isinstance(it, VizResult) else it
        flat.extend(fig if isinstance(fig, list) else [fig])
    titles = list(titles) if titles is not None else [None] * len(flat)
    prs = pptx.Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for fig, title in zip(flat, titles + [None] * (len(flat) - len(titles))):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        top = 0.3
        if title:
            box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12.3), Inches(0.7))
            box.text_frame.text = title
            box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
            top = 1.0
        buf = BytesIO()
        fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight")
        buf.seek(0)
        w, h = fig.get_size_inches()
        scale = min(12.3 / w, (7.5 - top - 0.3) / h)
        slide.shapes.add_picture(buf, Inches((13.333 - w * scale) / 2), Inches(top), Inches(w * scale), Inches(h * scale))
    prs.save(path)
    return str(path)
